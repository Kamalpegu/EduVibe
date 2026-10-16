"""
concanate.py - SVG slides + narration audio -> one PPTX (language selectable).

python concanate.py --lang en
python concanate.py --lang hi
python concanate.py --lang hi --no-autoplay
python concanate.py --lang en --svg-dir X --audio-dir Y --speech-json Z --out deck.pptx

pip install python-pptx playwright
playwright install chromium        (renders SVG -> PNG; or use cairosvg if it works for you)
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import tempfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

SRC = Path(__file__).resolve().parent.parent
OUT_DIR = Path(__file__).resolve().parent / "ppt_outputs"

LANGS = {
    "en": {
        "audio": SRC / "voice_over" / "audio",
        "speech": SRC / "voice_over" / "speech.json",
    },
    "hi": {
        "audio": SRC / "voice_over" / "audio_hi",
        "speech": SRC / "voice_over" / "speech_hi.json",
    },
}
DEFAULT_SVG_DIR = SRC / "image_create" / "slides_svg"


# ----------------------------------------------------------------------
# SVG -> PNG
# ----------------------------------------------------------------------

def _playwright(svgs: list[Path], tmp: Path) -> dict[Path, Path]:
    from playwright.sync_api import sync_playwright

    out: dict[Path, Path] = {}
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1280, "height": 720}, device_scale_factor=2)
        for s in svgs:
            png = tmp / (s.stem + ".png")
            page.goto(s.resolve().as_uri())
            page.screenshot(path=str(png))
            out[s] = png
        browser.close()
    return out


def _cairosvg(svgs: list[Path], tmp: Path) -> dict[Path, Path]:
    import cairosvg

    out: dict[Path, Path] = {}
    for s in svgs:
        png = tmp / (s.stem + ".png")
        cairosvg.svg2png(url=str(s), write_to=str(png), output_width=2560, output_height=1440)
        out[s] = png
    return out


def svgs_to_pngs(svgs: list[Path], tmp: Path, renderer: str = "auto") -> dict[Path, Path]:
    order = {"auto": [_playwright, _cairosvg], "playwright": [_playwright], "cairosvg": [_cairosvg]}[renderer]
    errors = []
    for fn in order:
        try:
            return fn(svgs, tmp)
        except Exception as e:
            errors.append(f"{fn.__name__}: {type(e).__name__}: {e}")
    sys.exit(
        "Could not convert SVG to PNG.\n  " + "\n  ".join(errors) +
        "\nFix: pip install playwright && playwright install chromium"
    )


# ----------------------------------------------------------------------
# Audio autoplay (PowerPoint timing node)
# ----------------------------------------------------------------------

def _autoplay_xml(spid: int) -> str:
    return f"""<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""


# ----------------------------------------------------------------------
# Build
# ----------------------------------------------------------------------

def _slide_no(svg: Path) -> int:
    m = re.match(r"slide_(\d+)", svg.stem)
    return int(m.group(1)) if m else -1


def _load_speech(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"deck_title": "", "slides": []}
    raw = re.sub(r"```(?:json)?", "", path.read_text(encoding="utf-8")).strip()
    return json.loads(raw)


def build_ppt(
    lang: str = "en",
    svg_dir: Path = DEFAULT_SVG_DIR,
    audio_dir: Path | None = None,
    speech_json: Path | None = None,
    out_path: Path | None = None,
    autoplay: bool = True,
    with_audio: bool = True,
    renderer: str = "auto",
) -> Path:
    if lang not in LANGS:
        raise ValueError(f"lang must be one of {list(LANGS)}")
    svg_dir = Path(svg_dir)
    audio_dir = Path(audio_dir or LANGS[lang]["audio"])
    speech_json = Path(speech_json or LANGS[lang]["speech"])

    svgs = sorted((p for p in svg_dir.glob("slide_*.svg") if _slide_no(p) >= 0), key=_slide_no)
    if not svgs:
        raise FileNotFoundError(f"No slide_*.svg files in {svg_dir}")

    speech = _load_speech(speech_json)
    narration = {s["slide_number"]: (s.get("speech") or "") for s in speech.get("slides", [])}
    title = speech.get("deck_title") or "presentation"

    if out_path is None:
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        safe = re.sub(r"[^\w\-]+", "_", title).strip("_") or "presentation"
        out_path = OUT_DIR / f"{safe}_{lang}.pptx"
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    missing = []
    with tempfile.TemporaryDirectory() as t:
        pngs = svgs_to_pngs(svgs, Path(t), renderer)

        for svg in svgs:
            n = _slide_no(svg)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(pngs[svg]), 0, 0, prs.slide_width, prs.slide_height)

            text = narration.get(n, "")
            if text:
                slide.notes_slide.notes_text_frame.text = text

            if with_audio:
                mp3 = audio_dir / f"slide_{n:02d}.mp3"
                if mp3.exists():
                    icon = slide.shapes.add_movie(
                        str(mp3), Inches(12.5), Inches(0.35), Inches(0.5), Inches(0.5),
                        mime_type="audio/mpeg",
                    )
                    if autoplay:
                        timing = slide._element.find(qn("p:timing"))
                        media_timing = (
                            timing.find(
                                ".//"
                                + qn("p:video")
                                + "/"
                                + qn("p:cMediaNode")
                                + "/"
                                + qn("p:cTn")
                                + "/"
                                + qn("p:stCondLst")
                            )
                            if timing is not None
                            else None
                        )
                        if media_timing is None:
                            raise RuntimeError("PowerPoint could not configure audio autoplay.")
                        for condition in media_timing.findall(qn("p:cond")):
                            if condition.get("delay") == "indefinite":
                                condition.set("delay", "0")
                else:
                    missing.append(mp3.name)

        prs.save(str(out_path))

    if missing:
        sys.stderr.write(f"[warn] no audio for: {', '.join(missing)} (looked in {audio_dir})\n")
    return out_path


def main() -> None:
    ap = argparse.ArgumentParser(description="SVG slides + audio -> PPTX")
    ap.add_argument("--lang", choices=list(LANGS), default="en", help="en or hi (picks audio + narration notes)")
    ap.add_argument("--svg-dir", type=Path, default=DEFAULT_SVG_DIR)
    ap.add_argument("--audio-dir", type=Path)
    ap.add_argument("--speech-json", type=Path)
    ap.add_argument("--out", type=Path, help="Output .pptx (default: ppt_outputs/<title>_<lang>.pptx)")
    ap.add_argument("--no-autoplay", action="store_true", help="Audio icon only; click to play")
    ap.add_argument("--no-audio", action="store_true")
    ap.add_argument("--renderer", choices=["auto", "playwright", "cairosvg"], default="auto")
    a = ap.parse_args()

    p = build_ppt(a.lang, a.svg_dir, a.audio_dir, a.speech_json, a.out,
                  autoplay=not a.no_autoplay, with_audio=not a.no_audio, renderer=a.renderer)
    print(f"Saved {p}")


if __name__ == "__main__":
    main()